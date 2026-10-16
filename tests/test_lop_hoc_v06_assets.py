# -*- coding: utf-8 -*-
"""Checks the fake v0.6 classroom assets (games + capability cards, Lớp học NNC).

Covers plugins/nnc/skills/work-first-learning/assets/mau/{gop,tim-loi,ban-la-sep},
phieu-nhan-hang-mau.png and thong-bao-noi-bo-mau.md, all produced by
_tao-du-lieu-v06.py. Skipped when openpyxl (or, for the image, Pillow) is missing.
"""
import datetime as dt
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None
try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

ROOT = Path(__file__).resolve().parents[1]
MAU = ROOT / "plugins" / "nnc" / "skills" / "work-first-learning" / "assets" / "mau"
GEN = MAU / "_tao-du-lieu-v06.py"
README = MAU / "README.md"
SUFFIX = "— mẫu"
KENH = ("OTC", "HC", "SPM")
COLS = ["Ngày", "Kênh", "Khách hàng", "Mã hàng", "Tên hàng",
        "Số lượng", "Đơn vị", "Đơn giá", "Thành tiền"]

GOP = {40: dt.date(2026, 9, 28), 41: dt.date(2026, 10, 5), 42: dt.date(2026, 10, 12)}
XLSX = ["gop/don-hang-tuan-40.xlsx", "gop/don-hang-tuan-41.xlsx", "gop/don-hang-tuan-42.xlsx",
        "tim-loi/bang-tim-loi.xlsx", "ban-la-sep/ket-qua-claude-mau.xlsx"]
TEXT = ["tim-loi/dap-an-tim-loi.md", "ban-la-sep/dap-an.md", "thong-bao-noi-bo-mau.md"]
PNG = "phieu-nhan-hang-mau.png"
NEW_FILES = XLSX + TEXT + [PNG, "_tao-du-lieu-v06.py"]

# Documented planted defects (README v0.6 + answer keys).
TIM_LOI_WEEK = (dt.date(2026, 10, 19), dt.date(2026, 10, 24))
TIM_LOI = {7: "kenh", 15: "thanh_tien", 21: "ngay"}
SEP_WRONG = ("HC", 3, -950000)          # channel, column index (Tổng Thành tiền), amount
NEW_CUSTOMER = "Nhà thuốc P — mẫu"
JUMP = ("MAU-HC-02", 20, 200)


def fmt(n):
    return format(n, ",").replace(",", ".")


def read_sheet(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    return rows


def read_orders(path):
    rows = read_sheet(path)
    header = list(rows[0])
    return header, [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def by_channel(data):
    out = {k: [0, 0, 0] for k in KENH}
    for r in data:
        t = out[r["Kênh"]]
        t[0] += 1
        t[1] += r["Số lượng"]
        t[2] += r["Thành tiền"]
    return out


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class V06Assets(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.readme = README.read_text(encoding="utf-8")
        cls.gop = {w: read_orders(MAU / ("gop/don-hang-tuan-%d.xlsx" % w)) for w in GOP}
        cls.tl_head, cls.tl = read_orders(MAU / "tim-loi/bang-tim-loi.xlsx")

    def test_files_exist(self):
        for name in NEW_FILES:
            self.assertTrue((MAU / name).is_file(), name)

    def test_readme_mentions_every_new_file_and_fake_notice(self):
        for name in NEW_FILES:
            self.assertIn("`%s`" % name, self.readme, name)
        self.assertIn("**TOÀN BỘ DỮ LIỆU TRONG THƯ MỤC NÀY LÀ GIẢ.**", self.readme)
        self.assertIn("**VẪN LÀ DỮ LIỆU GIẢ.**", self.readme)

    def test_all_names_are_marked_fake(self):
        for data in [d for _, d in self.gop.values()] + [self.tl]:
            for r in data:
                self.assertTrue(r["Khách hàng"].endswith(SUFFIX), r["Khách hàng"])
                self.assertTrue(r["Tên hàng"].endswith(SUFFIX), r["Tên hàng"])
                self.assertTrue(r["Mã hàng"].startswith("MAU-"), r["Mã hàng"])
        notice = (MAU / "thong-bao-noi-bo-mau.md").read_text(encoding="utf-8")
        self.assertIn("GIẢ", notice)
        self.assertIn("— mẫu", notice)
        self.assertTrue(80 <= len(notice.split()) <= 220, len(notice.split()))

    # --- gop ----------------------------------------------------------------
    def test_gop_consistent_columns_and_clean(self):
        for w, (head, data) in self.gop.items():
            self.assertEqual(head, COLS, w)
            self.assertTrue(18 <= len(data) <= 22, (w, len(data)))
            keys = [(r["Khách hàng"], r["Mã hàng"]) for r in data]
            self.assertEqual(len(keys), len(set(keys)), w)
            for r in data:
                self.assertIn(r["Kênh"], KENH)
                self.assertEqual(r["Đơn vị"], "hộp")
                self.assertEqual(r["Thành tiền"], r["Số lượng"] * r["Đơn giá"])
                self.assertTrue(GOP[w] <= r["Ngày"].date() <= GOP[w] + dt.timedelta(days=5), (w, r))

    def test_gop_week41_vs_42_three_explainable_changes(self):
        q41 = {(r["Khách hàng"], r["Mã hàng"]): r for r in self.gop[41][1]}
        q42 = {(r["Khách hàng"], r["Mã hàng"]): r for r in self.gop[42][1]}
        added, removed = set(q42) - set(q41), set(q41) - set(q42)
        changed = {k for k in set(q41) & set(q42) if q41[k]["Số lượng"] != q42[k]["Số lượng"]}
        self.assertEqual({k[0] for k in added}, {NEW_CUSTOMER})
        self.assertEqual(len(added), 2)
        self.assertNotIn(NEW_CUSTOMER, {r["Khách hàng"] for w in (40, 41) for r in self.gop[w][1]})
        self.assertEqual(len(removed), 2)
        self.assertEqual({q41[k]["Kênh"] for k in removed}, {"SPM"})
        self.assertEqual(len(changed), 1)
        (k,) = changed
        self.assertEqual((k[1], q41[k]["Số lượng"], q42[k]["Số lượng"]), JUMP)
        t41, t42 = by_channel(self.gop[41][1]), by_channel(self.gop[42][1])
        self.assertGreater(t42["OTC"][2], t41["OTC"][2])
        self.assertGreater(t42["HC"][2], t41["HC"][2])
        self.assertLess(t42["SPM"][2], t41["SPM"][2])
        # Row numbers documented in README.
        rows42 = [n for n, r in enumerate(self.gop[42][1], start=2) if r["Khách hàng"] == NEW_CUSTOMER]
        self.assertIn("Dòng %d và %d: **%s**" % (rows42[0], rows42[1], NEW_CUSTOMER), self.readme)
        for w in (41, 42):
            n = [n for n, r in enumerate(self.gop[w][1], start=2) if r["Mã hàng"] == JUMP[0]][0]
            self.assertIn("Dòng %d:" % n, self.readme)
        for n, r in enumerate(self.gop[41][1], start=2):
            if (r["Khách hàng"], r["Mã hàng"]) in removed:
                self.assertIn("ng %d: %s · %s" % (n, r["Khách hàng"], r["Tên hàng"]), self.readme)
        for kenh in KENH:
            self.assertIn(fmt(abs(t42[kenh][2] - t41[kenh][2])), self.readme, kenh)

    def test_gop_readme_totals_match(self):
        grand = {k: [0, 0, 0] for k in KENH}
        for w in GOP:
            t = by_channel(self.gop[w][1])
            for k in KENH:
                grand[k] = [a + b for a, b in zip(grand[k], t[k])]
            line = "| Tuần %d | %d | %s | %s | %s | %s | %s |" % (
                w, len(self.gop[w][1]), fmt(t["OTC"][2]), fmt(t["HC"][2]), fmt(t["SPM"][2]),
                fmt(sum(t[k][1] for k in KENH)), fmt(sum(t[k][2] for k in KENH)))
            self.assertIn(line, self.readme)
        line = "| **Gộp 3 tuần** | **%d** | **%s** | **%s** | **%s** | **%s** | **%s** |" % (
            sum(grand[k][0] for k in KENH), fmt(grand["OTC"][2]), fmt(grand["HC"][2]), fmt(grand["SPM"][2]),
            fmt(sum(grand[k][1] for k in KENH)), fmt(sum(grand[k][2] for k in KENH)))
        self.assertIn(line, self.readme)

    # --- tim-loi ------------------------------------------------------------
    def test_tim_loi_exactly_three_planted_errors(self):
        self.assertEqual(self.tl_head, COLS)
        self.assertTrue(23 <= len(self.tl) <= 27, len(self.tl))
        found = {}
        for n, r in enumerate(self.tl, start=2):
            kinds = []
            if r["Kênh"] not in KENH:
                kinds.append("kenh")
            if r["Thành tiền"] != r["Số lượng"] * r["Đơn giá"]:
                kinds.append("thanh_tien")
            if not TIM_LOI_WEEK[0] <= r["Ngày"].date() <= TIM_LOI_WEEK[1]:
                kinds.append("ngay")
            self.assertEqual(r["Đơn vị"], "hộp", n)
            for c in COLS:
                self.assertIsNotNone(r[c], (n, c))
            if kinds:
                self.assertEqual(len(kinds), 1, (n, kinds))
                found[n] = kinds[0]
        self.assertEqual(found, TIM_LOI)
        self.assertEqual(self.tl[7 - 2]["Kênh"], "OCT")
        keys = [(r["Khách hàng"], r["Mã hàng"]) for r in self.tl]
        self.assertEqual(len(keys), len(set(keys)))
        key = (MAU / "tim-loi/dap-an-tim-loi.md").read_text(encoding="utf-8")
        for n in TIM_LOI:
            self.assertIn("| %d |" % n, key)
            self.assertIn("| %d |" % n, self.readme)

    # --- ban-la-sep ---------------------------------------------------------
    def test_ban_la_sep_one_wrong_total(self):
        _, src = read_orders(MAU / "bang-tinh-mau.xlsx")
        calc = by_channel(src)
        grand = [sum(calc[k][i] for k in KENH) for i in range(3)]
        rows = read_sheet(MAU / "ban-la-sep/ket-qua-claude-mau.xlsx")
        self.assertEqual(list(rows[0]), ["Kênh", "Số dòng", "Tổng Số lượng (hộp)", "Tổng Thành tiền (LAK)"])
        shown = {r[0]: list(r[1:4]) for r in rows[1:5]}
        expected = dict(calc, **{"Tổng cộng": grand})
        diffs = [(k, i + 1, shown[k][i] - expected[k][i])
                 for k in expected for i in range(3) if shown[k][i] != expected[k][i]]
        self.assertEqual(diffs, [SEP_WRONG])
        key = (MAU / "ban-la-sep/dap-an.md").read_text(encoding="utf-8")
        self.assertIn("lệch **%s LAK**" % fmt(SEP_WRONG[2]), key)
        self.assertIn("Ô **D3**", key)
        self.assertIn(fmt(calc["HC"][2]), key)
        self.assertIn(fmt(calc["HC"][2] + SEP_WRONG[2]), self.readme)
        self.assertIn("−%s LAK" % fmt(-SEP_WRONG[2]), self.readme)
        # Missing note: the sheet never states its source file or period.
        text = " ".join(str(v) for r in rows for v in r if v is not None)
        self.assertNotIn("bang-tinh-mau", text)
        self.assertNotIn("Nguồn", text)
        self.assertIn("Ghi chú", text)
        self.assertIn("nguồn và kỳ dữ liệu", key)

    # --- image ----------------------------------------------------------------
    @unittest.skipIf(Image is None, "Pillow not installed")
    def test_png_size(self):
        with Image.open(MAU / PNG) as im:
            self.assertEqual(im.format, "PNG")
            self.assertEqual(im.size, (1000, 700))

    # --- generator ------------------------------------------------------------
    def test_generator_is_deterministic(self):
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run([sys.executable, str(GEN), "--out", tmp], check=True,
                           stdout=subprocess.DEVNULL)
            for name in XLSX:
                a = openpyxl.load_workbook(MAU / name)
                b = openpyxl.load_workbook(Path(tmp) / name)
                self.assertEqual(a.sheetnames, b.sheetnames, name)
                for s in a.sheetnames:
                    self.assertEqual(list(a[s].iter_rows(values_only=True)),
                                     list(b[s].iter_rows(values_only=True)), (name, s))
                a.close()
                b.close()
            for name in TEXT:
                self.assertEqual((MAU / name).read_text(encoding="utf-8"),
                                 (Path(tmp) / name).read_text(encoding="utf-8"), name)
            if Image is not None:
                with Image.open(MAU / PNG) as a, Image.open(Path(tmp) / PNG) as b:
                    self.assertEqual(a.size, b.size)
                    self.assertEqual(a.tobytes(), b.tobytes())


if __name__ == "__main__":
    unittest.main()
