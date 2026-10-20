# -*- coding: utf-8 -*-
"""Checks the fake classroom sample data (Lớp học NNC, Buổi 1).

Verifies the xlsx files under plugins/nnc/skills/work-first-learning/assets/mau
are consistent with their answer key and contain exactly the planted defects.
Skipped when openpyxl is not installed.
"""
from collections import Counter
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

ROOT = Path(__file__).resolve().parents[1]
WFL = ROOT / "plugins" / "nnc" / "skills" / "work-first-learning"
MAU = WFL / "assets" / "mau"
GEN = MAU / "_tao-du-lieu.py"
SKILL_MAU = WFL / "assets" / "skill-mau" / "tong-hop-don-hang" / "SKILL.md"
SUFFIX = "— mẫu"
KENH = ("OTC", "HC", "SPM")
COLS = ["Ngày", "Kênh", "Khách hàng", "Mã hàng", "Tên hàng",
        "Số lượng", "Đơn vị", "Đơn giá", "Thành tiền"]


def read_sheet(path, sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def read_orders(path):
    rows = read_sheet(path)
    header = list(rows[0])
    data = [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]
    return header, data


def recompute(data, skip_rows=()):
    out = {k: [0, 0, 0] for k in KENH}
    for n, r in enumerate(data, start=2):
        if n in skip_rows:
            continue
        t = out[r["Kênh"]]
        t[0] += 1
        t[1] += r["Số lượng"]
        t[2] += r["Thành tiền"]
    return out


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class SampleData(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.h1, cls.w1 = read_orders(MAU / "bang-tinh-mau.xlsx")
        cls.h2, cls.w2 = read_orders(MAU / "bang-tinh-mau-tuan-sau.xlsx")
        cls.key = read_sheet(MAU / "dap-an-mau.xlsx", "Tổng hợp theo kênh")
        cls.defects = [r for r in read_sheet(MAU / "dap-an-mau.xlsx", "Lỗi cài sẵn tuần sau")[1:]
                       if r[0] and isinstance(r[1], int)]
        cls.expect2 = read_sheet(MAU / "dap-an-mau.xlsx", "Tuần sau — kết quả mong đợi")

    def defect_rows(self, kind):
        return sorted(r[1] for r in self.defects if r[0] == kind)

    # --- bang-tinh-mau.xlsx -------------------------------------------------
    def test_week1_shape(self):
        self.assertEqual(self.h1, COLS)
        self.assertTrue(38 <= len(self.w1) <= 42, len(self.w1))
        self.assertEqual({r["Kênh"] for r in self.w1}, set(KENH))

    def test_week1_clean(self):
        keys = Counter((r["Mã hàng"], r["Khách hàng"], r["Ngày"], r["Số lượng"]) for r in self.w1)
        self.assertEqual([k for k, c in keys.items() if c > 1], [])
        for r in self.w1:
            self.assertEqual(r["Đơn vị"], "hộp")
            self.assertIsInstance(r["Đơn giá"], int)
            self.assertEqual(r["Đơn giá"] % 1000, 0)
            self.assertEqual(r["Thành tiền"], r["Số lượng"] * r["Đơn giá"])
            self.assertEqual((r["Ngày"].year, r["Ngày"].month), (2026, 10))
        days = {r["Ngày"].date() for r in self.w1}
        self.assertLessEqual((max(days) - min(days)).days, 6)

    def test_all_names_are_marked_fake(self):
        for data in (self.w1, self.w2):
            for r in data:
                self.assertTrue(r["Khách hàng"].endswith(SUFFIX), r["Khách hàng"])
                self.assertTrue(r["Tên hàng"].endswith(SUFFIX), r["Tên hàng"])
                self.assertTrue(r["Mã hàng"].startswith("MAU-"), r["Mã hàng"])

    def test_answer_key_matches_recomputation(self):
        calc = recompute(self.w1)
        rows = {r[0]: r[1:4] for r in self.key[1:] if r[0]}
        for k in KENH:
            self.assertEqual(list(rows[k]), calc[k], k)
        grand = [sum(calc[k][i] for k in KENH) for i in range(3)]
        self.assertEqual(list(rows["Tổng cộng"]), grand)
        self.assertEqual(grand[0], len(self.w1))

    # --- bang-tinh-mau-tuan-sau.xlsx ---------------------------------------
    def test_week2_renamed_column(self):
        self.assertEqual(self.h2, [("Giá bán" if c == "Đơn giá" else c) for c in COLS])
        self.assertEqual(self.defect_rows("Cột đổi tên"), [1])

    def test_week2_exactly_three_missing_prices(self):
        missing = [n for n, r in enumerate(self.w2, start=2) if r["Giá bán"] is None]
        self.assertEqual(len(missing), 3)
        self.assertEqual(missing, self.defect_rows("Thiếu giá"))
        self.assertEqual({self.w2[n - 2]["Kênh"] for n in missing}, set(KENH))
        for n, r in enumerate(self.w2, start=2):
            for c in ("Ngày", "Kênh", "Khách hàng", "Mã hàng", "Tên hàng", "Số lượng", "Đơn vị"):
                self.assertIsNotNone(r[c], (n, c))
            if n not in missing:
                self.assertEqual(r["Thành tiền"], r["Số lượng"] * r["Giá bán"], n)

    def test_week2_exactly_one_duplicate(self):
        seen, dups = {}, []
        for n, r in enumerate(self.w2, start=2):
            k = (r["Mã hàng"], r["Khách hàng"], r["Ngày"], r["Số lượng"])
            if k in seen:
                dups.append((seen[k], n))
            else:
                seen[k] = n
        self.assertEqual(len(dups), 1, dups)
        self.assertEqual(self.defect_rows("Dòng trùng"), [dups[0][1]])

    def test_week2_exactly_one_unit_mismatch(self):
        units = Counter(r["Đơn vị"] for r in self.w2)
        self.assertEqual(units, Counter({"hộp": len(self.w2) - 1, "vỉ": 1}))
        vi = [n for n, r in enumerate(self.w2, start=2) if r["Đơn vị"] == "vỉ"]
        self.assertEqual(vi, self.defect_rows("Đơn vị lệch"))

    def test_week2_defects_do_not_overlap_and_nothing_else_listed(self):
        self.assertEqual(Counter(r[0] for r in self.defects),
                         Counter({"Cột đổi tên": 1, "Thiếu giá": 3, "Dòng trùng": 1, "Đơn vị lệch": 1}))
        data_rows = [r[1] for r in self.defects if r[0] != "Cột đổi tên"]
        self.assertEqual(len(set(data_rows)), 5)
        self.assertEqual(len(self.w2), len(self.w1) + 1)

    def test_week2_expected_skill_result(self):
        flagged = {r[1] for r in self.defects if r[0] != "Cột đổi tên"}
        data = [dict(r, **{"Thành tiền": r["Thành tiền"] or 0}) for r in self.w2]
        calc = recompute(data, skip_rows=flagged)
        rows = {r[0]: r[1:5] for r in self.expect2[1:] if r[0]}
        for k in KENH:
            n_flag = sum(1 for n in flagged if self.w2[n - 2]["Kênh"] == k)
            self.assertEqual(list(rows[k]), calc[k] + [n_flag], k)
        self.assertEqual(rows["Tổng cộng"][3], 5)

    # --- docs and generator ------------------------------------------------
    def test_readme_and_skill_mention_planted_rows(self):
        readme = (MAU / "README.md").read_text(encoding="utf-8")
        skill = SKILL_MAU.read_text(encoding="utf-8")
        self.assertIn("**TOÀN BỘ DỮ LIỆU TRONG THƯ MỤC NÀY LÀ GIẢ.**", readme)
        self.assertTrue(skill.startswith("---\nname: tong-hop-don-hang\n"))
        self.assertIn("Cần kiểm", skill)
        for r in self.defects:
            if r[0] != "Cột đổi tên":
                self.assertIn("| %d |" % r[1], readme, r)
                self.assertIn(str(r[1]), skill, r)

    def test_generator_is_deterministic(self):
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run([sys.executable, str(GEN), "--out", tmp], check=True,
                           stdout=subprocess.DEVNULL)
            for name in ("bang-tinh-mau.xlsx", "bang-tinh-mau-tuan-sau.xlsx", "dap-an-mau.xlsx"):
                a = openpyxl.load_workbook(MAU / name)
                b = openpyxl.load_workbook(Path(tmp) / name)
                self.assertEqual(a.sheetnames, b.sheetnames, name)
                for s in a.sheetnames:
                    self.assertEqual(list(a[s].iter_rows(values_only=True)),
                                     list(b[s].iter_rows(values_only=True)), (name, s))
                a.close()
                b.close()


if __name__ == "__main__":
    unittest.main()
